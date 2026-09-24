"""Read-only XLSX -> private JSONL evidence, one record per physical data row.

The script never prints cell values or file hashes. It does not evaluate formulas,
merge duplicate clients, or write to a database.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import pathlib
import sys
import tempfile
import zipfile
from collections import Counter

from openpyxl import load_workbook

MAX_SOURCE_BYTES = 256 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_ZIP_PARTS = 1000
MAX_ROWS = 200_000


class ExtractError(Exception):
    pass


def digest(path: pathlib.Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def check_xlsx(path: pathlib.Path) -> None:
    if not path.is_file() or path.stat().st_size > MAX_SOURCE_BYTES:
        raise ExtractError("SOURCE_SIZE_OR_TYPE")
    with zipfile.ZipFile(path) as archive:
        parts = archive.infolist()
        if len(parts) > MAX_ZIP_PARTS or sum(part.file_size for part in parts) > MAX_UNCOMPRESSED_BYTES:
            raise ExtractError("XLSX_SIZE_LIMIT")
        for part in parts:
            if part.file_size > MAX_SOURCE_BYTES or part.flag_bits & 0x1:
                raise ExtractError("XLSX_UNSAFE_PART")
            if part.filename.lower().endswith((".xml", ".rels")):
                # Reject entity declarations before openpyxl reads the workbook.
                tail = b""
                with archive.open(part) as xml:
                    for chunk in iter(lambda: xml.read(64 * 1024), b""):
                        probe = (tail + chunk).upper()
                        if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
                            raise ExtractError("XLSX_XML_ENTITIES")
                        tail = probe[-16:]


def cell_evidence(cell: object) -> dict[str, str | None]:
    value = getattr(cell, "value", None)
    kind = getattr(cell, "data_type", None)
    number_format = getattr(cell, "number_format", None)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        lexical = value.isoformat()
    elif value is None:
        lexical = None
    else:
        lexical = str(value)
    return {"type": kind, "value": lexical, "number_format": number_format}


def write_private(path: pathlib.Path, content: bytes) -> None:
    if path.exists():
        if hashlib.sha256(path.read_bytes()).digest() != hashlib.sha256(content).digest():
            raise ExtractError("OUTPUT_ALREADY_EXISTS_DIFFERENT")
        return
    descriptor, temporary = tempfile.mkstemp(prefix=".kc-evidence-", dir=path.parent)
    try:
        os.chmod(temporary, 0o600)
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        if path.exists():
            raise ExtractError("OUTPUT_ALREADY_EXISTS_DIFFERENT")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main() -> None:
    parser = argparse.ArgumentParser(add_help=True)
    parser.add_argument("--dataset", required=True, choices=("clients", "expenses"))
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--expected-rows", type=int)
    parser.add_argument("--expected-columns", type=int)
    options = parser.parse_args()
    source = pathlib.Path(options.input).resolve(strict=True)
    target = pathlib.Path(options.output).resolve()
    manifest_path = pathlib.Path(options.manifest).resolve()
    if source == target or source == manifest_path or target == manifest_path:
        raise ExtractError("PATH_COLLISION")
    if target.parent != source.parent or manifest_path.parent != source.parent:
        raise ExtractError("OUTPUT_MUST_STAY_IN_PRIVATE_SOURCE_DIRECTORY")
    check_xlsx(source)
    source_hash = digest(source)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ExtractError("ONE_WORKSHEET_REQUIRED")
        sheet = workbook.worksheets[0]
        if options.expected_columns is not None and sheet.max_column != options.expected_columns:
            raise ExtractError("COLUMN_COUNT_MISMATCH")
        if sheet.max_column is None or sheet.max_column < 1 or sheet.max_column > 100:
            raise ExtractError("COLUMN_COUNT_LIMIT")
        rows = sheet.iter_rows(min_row=1, max_col=sheet.max_column)
        header = next(rows, None)
        if header is None:
            raise ExtractError("EMPTY_WORKSHEET")
        labels = [cell_evidence(cell) for cell in header]
        lines: list[bytes] = []
        duplicate_counts: Counter[str] = Counter()
        count = 0
        for physical_row, cells in enumerate(rows, start=2):
            count += 1
            if count > MAX_ROWS:
                raise ExtractError("ROW_COUNT_LIMIT")
            evidence_cells = [cell_evidence(cell) for cell in cells]
            row_id = hashlib.sha256(f"{source_hash}\0{sheet.title}\0{physical_row}".encode()).hexdigest()
            record = {
                "evidence_row_id": row_id,
                "source_file_sha256": source_hash,
                "source_worksheet": sheet.title,
                "source_row_number": physical_row,
                "source_columns": labels,
                "cells": evidence_cells,
            }
            lines.append((json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8"))
            duplicate_counts[hashlib.sha256(json.dumps(evidence_cells, ensure_ascii=False, sort_keys=True).encode()).hexdigest()] += 1
        if options.expected_rows is not None and count != options.expected_rows:
            raise ExtractError("ROW_COUNT_MISMATCH")
        manifest = {
            "dataset": f"realtycalendar-{options.dataset}-v1",
            "sha256": source_hash,
            "worksheet_names": [sheet.title],
            "row_count_by_sheet": {sheet.title: count},
            "column_count_by_sheet": {sheet.title: sheet.max_column},
            "parser_version": "extract_xlsx_evidence_v1",
        }
        # Files are private artifacts beside the original. Existing files are
        # accepted only if byte-identical, so reruns cannot silently rewrite a snapshot.
        write_private(target, b"".join(lines))
        write_private(manifest_path, json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
        duplicates = sum(count - 1 for count in duplicate_counts.values())
        print(json.dumps({"dataset": options.dataset, "rows": count, "columns": sheet.max_column, "duplicate_rows_preserved": duplicates}))
    finally:
        workbook.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        # openpyxl exceptions can contain a source value; terminal logs only
        # receive stable codes, never workbook cells or private paths.
        code = str(error) if isinstance(error, ExtractError) else "EXTRACTION_FAILED"
        print(f"XLSX evidence extraction failed: {code}", file=sys.stderr)
        sys.exit(1)
